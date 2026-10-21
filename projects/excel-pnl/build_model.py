"""
Собирает Excel-модель «P&L магазина на маркетплейсах» (pnl_model.xlsx).
Все расчёты в файле — формулы Excel: меняете допущения на листе «Допущения», модель пересчитывается.

Запуск: python projects/excel-pnl/build_model.py
"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import ColorScaleRule, CellIsRule
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.comments import Comment

OUT = Path(__file__).resolve().parent / "pnl_model.xlsx"
wb = Workbook()

# ---------- стили ----------
F = "Arial"
f_base = Font(name=F, size=10)
f_in = Font(name=F, size=10, color="0000FF")          # ввод — синий
f_link = Font(name=F, size=10, color="008000")        # ссылка на другой лист — зелёный
f_b = Font(name=F, size=10, bold=True)
f_h1 = Font(name=F, size=14, bold=True)
f_h2 = Font(name=F, size=11, bold=True, color="FFFFFF")
f_note = Font(name=F, size=9, italic=True, color="666666")
fill_key = PatternFill("solid", fgColor="FFFF00")
fill_head = PatternFill("solid", fgColor="2F5FA6")
fill_sub = PatternFill("solid", fgColor="E8EEF7")
fill_tot = PatternFill("solid", fgColor="F2F2F2")
thin = Side(style="thin", color="BFBFBF")
b_bot = Border(bottom=thin)
b_top = Border(top=Side(style="thin", color="000000"))
RUB = '#,##0" ₽";(#,##0" ₽");"-"'
RUB2 = '#,##0.0" ₽";(#,##0.0" ₽");"-"'
PCT = '0.0%;(0.0%);"-"'
NUM = '#,##0;(#,##0);"-"'
IDX = '0.00'

def cell(ws, ref, v, font=f_base, fmt=None, fill=None, al=None, border=None, comment=None):
    c = ws[ref]; c.value = v; c.font = font
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    if al: c.alignment = al
    if border: c.border = border
    if comment: c.comment = Comment(comment, "Модель")
    return c

def header_row(ws, row, labels, start=1):
    for i, t in enumerate(labels):
        cell(ws, f"{L(start + i)}{row}", t, f_h2, fill=fill_head, al=Alignment(horizontal="center", vertical="center", wrap_text=True))
    ws.row_dimensions[row].height = 30

CH = ["Wildberries", "Ozon", "Сайт"]
MONTHS = ["янв", "фев", "мар", "апр", "май", "июн", "июл", "авг", "сен", "окт", "ноя", "дек"]

# =====================================================================
# 1. Инструкция
# =====================================================================
ws = wb.active; ws.title = "Инструкция"
ws.column_dimensions["A"].width = 3; ws.column_dimensions["B"].width = 100
cell(ws, "B2", "P&L магазина товаров для дома на маркетплейсах — модель на 2027 год", f_h1)
lines = [
    ("Зачем", f_b),
    ("Собственник спрашивает: сколько мы заработаем за год, какие товары и каналы убыточны и что будет, если поднять цену или урезать рекламу.", f_base),
    ("", f_base),
    ("Как устроено", f_b),
    ("1. «Допущения» — комиссии, логистика, выкуп, ДРР по каналам, налог, постоянные расходы, сезонность и сценарии.", f_base),
    ("2. «Товары» — ассортимент: цена, себестоимость, план заказов в месяц и распределение по каналам.", f_base),
    ("3. «Юнит-экономика» — прибыль с одной проданной единицы по каждому товару в каждом канале и предельный ДРР.", f_base),
    ("4. «P&L» — помесячный отчёт о прибылях и убытках с учётом сезонности и роста, итоги по каналам и категориям, график.", f_base),
    ("5. «Чувствительность» — годовая операционная прибыль при разной цене и доле рекламы (таблица «что, если»).", f_base),
    ("", f_base),
    ("Обозначения", f_b),
    ("Синий шрифт — вводимые значения, их можно менять.   Чёрный — формулы, не трогать.   Зелёный — ссылка на другой лист.", f_base),
    ("Жёлтая заливка — ключевые допущения и переключатель сценария (лист «Допущения», ячейка C22).", f_base),
    ("", f_base),
    ("Попробуйте", f_b),
    ("• Переключите сценарий на «Пессимистичный» и посмотрите, сколько месяцев модель уходит в минус.", f_base),
    ("• Поднимите ДРР Wildberries до 20% и найдите товары, которые стали убыточными (красные строки на «Юнит-экономике»).", f_base),
    ("• Посмотрите на «Чувствительности», что сильнее влияет на прибыль: +5% к цене или −5 п.п. рекламы.", f_base),
    ("", f_base),
    ("Ограничения", f_b),
    ("Данные демонстрационные. Спрос не зависит от цены (эластичность не моделируется). Логистика — средняя за единицу, без учёта габаритов и складов.", f_note),
]
for i, (t, fnt) in enumerate(lines, start=4):
    cell(ws, f"B{i}", t, fnt, al=Alignment(wrap_text=True, vertical="top"))

# =====================================================================
# 2. Допущения
# =====================================================================
a = wb.create_sheet("Допущения")
for col, w in zip("ABCDEFGHIJKLMN", [3, 26, 14, 14, 14, 14, 14, 14, 10, 10, 10, 10, 10, 10]):
    a.column_dimensions[col].width = w
cell(a, "B2", "Допущения", f_h1)
cell(a, "B3", "Синие ячейки можно менять. Источник цифр по умолчанию — типичные тарифы и показатели продавца товаров для дома (демонстрационные).", f_note)

header_row(a, 5, ["Канал", "Комиссия, % от цены", "Логистика до покупателя, ₽/шт", "Обратная логистика, ₽/шт", "Выкуп, %", "ДРР, % от цены", "Хранение, ₽/шт"], start=2)
chan = [("Wildberries", 0.24, 92, 50, 0.78, 0.12, 18), ("Ozon", 0.21, 85, 45, 0.85, 0.10, 15), ("Сайт", 0.035, 160, 160, 0.95, 0.13, 0)]
for r, (n, com, log, back, buy, drr, sto) in enumerate(chan, start=6):
    cell(a, f"B{r}", n, f_b)
    for col, v, fmt in [("C", com, PCT), ("D", log, RUB), ("E", back, RUB), ("F", buy, PCT), ("G", drr, PCT), ("H", sto, RUB)]:
        cell(a, f"{col}{r}", v, f_in, fmt, fill=fill_key if col in "FG" else None)
a["C8"].comment = Comment("Для сайта — эквайринг. Логистика сайта — курьер/СДЭК за счёт магазина.", "Модель")
a["F6"].comment = Comment("Доля заказов, которые покупатель выкупил. Невыкупленный товар едет обратно за счёт продавца.", "Модель")

cell(a, "B11", "Налог (УСН «доходы»), % от выручки", f_base); cell(a, "C11", 0.06, f_in, PCT)
cell(a, "B12", "Постоянные расходы, ₽ в месяц", f_base); cell(a, "C12", 450000, f_in, RUB, fill=fill_key,
     comment="ФОТ 2 сотрудников, аренда склада, сервисы аналитики, бухгалтерия.")
cell(a, "B13", "Рост продаж, % в месяц", f_base); cell(a, "C13", 0.02, f_in, PCT)

header_row(a, 16, ["Сценарий", "Изменение цены", "Множитель ДРР", "Изменение выкупа, п.п.", "Изменение спроса"], start=2)
scen = [("Базовый", 0, 1.0, 0, 0), ("Оптимистичный", 0.05, 0.85, 0.03, 0.10), ("Пессимистичный", -0.05, 1.30, -0.05, -0.15)]
for r, (n, dp, dm, db, dv) in enumerate(scen, start=17):
    cell(a, f"B{r}", n, f_b)
    for col, v, fmt in [("C", dp, PCT), ("D", dm, IDX), ("E", db, PCT), ("F", dv, PCT)]:
        cell(a, f"{col}{r}", v, f_in, fmt)

cell(a, "B22", "Выбранный сценарий  ▸", f_b); cell(a, "C22", "Базовый", f_in, fill=fill_key, al=Alignment(horizontal="center"))
dv = DataValidation(type="list", formula1="=$B$17:$B$19", allow_blank=False)
a.add_data_validation(dv); dv.add("C22")
for r, (lbl, col, fmt) in enumerate([("Изменение цены", "C", PCT), ("Множитель ДРР", "D", IDX), ("Изменение выкупа, п.п.", "E", PCT), ("Изменение спроса", "F", PCT)], start=23):
    cell(a, f"B{r}", f"  {lbl}", f_base)
    cell(a, f"C{r}", f"=INDEX(${col}$17:${col}$19,MATCH($C$22,$B$17:$B$19,0))", f_base, fmt)

cell(a, "B29", "Сезонность (индекс месяца, 1,00 = средний месяц)", f_b)
season = [0.92, 0.88, 1.08, 0.93, 0.94, 0.86, 0.87, 0.97, 1.02, 1.00, 1.13, 1.40]
for i, (m, v) in enumerate(zip(MONTHS, season)):
    cell(a, f"{L(3 + i)}30", m, f_b, al=Alignment(horizontal="center"), fill=fill_sub)
    cell(a, f"{L(3 + i)}31", v, f_in, IDX, al=Alignment(horizontal="center"))
cell(a, "B31", "Индекс", f_base)
cell(a, "B32", "Проверка: среднее индексов", f_note); cell(a, "C32", "=AVERAGE(C31:N31)", f_note, IDX)
a.freeze_panes = "A5"

# =====================================================================
# 3. Товары
# =====================================================================
t = wb.create_sheet("Товары")
for col, w in zip("ABCDEFGHIJ", [3, 10, 30, 14, 13, 13, 15, 11, 11, 11]):
    t.column_dimensions[col].width = w
cell(t, "B2", "Ассортимент и план заказов", f_h1)
cell(t, "B3", "Заказы — до выкупа. Доля сайта считается как остаток. Цены и себестоимость демонстрационные.", f_note)
header_row(t, 5, ["SKU", "Товар", "Категория", "Цена, ₽", "Себестоимость, ₽", "Заказов в месяц, шт", "Доля WB", "Доля Ozon", "Доля сайта"], start=2)
goods = [
    ("DU-0001", "Плед вязаный", "Текстиль", 2490, 980, 420, 0.55, 0.30),
    ("DU-0002", "Комплект постельного белья", "Текстиль", 3990, 1650, 300, 0.50, 0.30),
    ("DU-0003", "Полотенце махровое", "Текстиль", 890, 340, 650, 0.50, 0.35),
    ("DU-0004", "Кружка керамическая", "Посуда", 690, 230, 900, 0.45, 0.35),
    ("DU-0005", "Набор тарелок", "Посуда", 2890, 1150, 260, 0.40, 0.35),
    ("DU-0006", "Френч-пресс", "Посуда", 1490, 560, 380, 0.45, 0.35),
    ("DU-0007", "Органайзер для ящика", "Хранение", 590, 260, 520, 0.55, 0.30),
    ("DU-0008", "Корзина плетёная", "Хранение", 690, 380, 480, 0.55, 0.35),
    ("DU-0009", "Свеча ароматическая", "Декор", 490, 210, 1100, 0.50, 0.35),
    ("DU-0010", "Ваза стеклянная", "Декор", 1890, 720, 240, 0.40, 0.35),
    ("DU-0011", "Коврик для ванной", "Ванная", 1290, 520, 360, 0.50, 0.35),
    ("DU-0012", "Держатель для щёток", "Ванная", 390, 190, 700, 0.55, 0.35),
]
G0, G1 = 6, 6 + len(goods) - 1
for r, (sku, name, cat, price, cost, orders, s_wb, s_oz) in enumerate(goods, start=G0):
    cell(t, f"B{r}", sku, f_in); cell(t, f"C{r}", name, f_in); cell(t, f"D{r}", cat, f_in)
    cell(t, f"E{r}", price, f_in, RUB); cell(t, f"F{r}", cost, f_in, RUB); cell(t, f"G{r}", orders, f_in, NUM)
    cell(t, f"H{r}", s_wb, f_in, PCT); cell(t, f"I{r}", s_oz, f_in, PCT)
    cell(t, f"J{r}", f"=1-H{r}-I{r}", f_base, PCT)
tr = G1 + 1
cell(t, f"C{tr}", "Итого / средневзвешенно", f_b, fill=fill_tot)
for col in "BDEFHIJ":
    t[f"{col}{tr}"].fill = fill_tot
cell(t, f"G{tr}", f"=SUM(G{G0}:G{G1})", f_b, NUM, fill=fill_tot)
cell(t, f"E{tr}", f"=SUMPRODUCT(E{G0}:E{G1},G{G0}:G{G1})/G{tr}", f_b, RUB, fill=fill_tot)
cell(t, f"F{tr}", f"=SUMPRODUCT(F{G0}:F{G1},G{G0}:G{G1})/G{tr}", f_b, RUB, fill=fill_tot)
t.conditional_formatting.add(f"J{G0}:J{G1}", CellIsRule(operator="lessThan", formula=["0"], font=Font(name=F, color="C00000", bold=True)))
t.freeze_panes = "A6"

# =====================================================================
# 4. Юнит-экономика (SKU × канал)
# =====================================================================
u = wb.create_sheet("Юнит-экономика")
cols = ["SKU", "Канал", "Товар", "Категория", "Цена, ₽", "Себест., ₽", "Комиссия, ₽", "Логистика на 1 выкуп, ₽", "Хранение, ₽", "Реклама, ₽", "Налог, ₽",
        "Прибыль с 1 шт, ₽", "Маржин-сть, %", "Предельный ДРР, %", "Заказов в мес", "Выкуп, %", "Продаж в мес", "Выручка в мес, ₽", "Прибыль в мес, ₽",
        "Базовая цена, ₽", "Комиссия, %", "Логистика, ₽/шт", "Обратная, ₽/шт"]
widths = [10, 12, 26, 11, 10, 10, 10, 13, 9, 10, 9, 11, 10, 11, 10, 8, 10, 13, 13, 10, 10, 10, 10]
for i, w in enumerate(widths, start=1):
    u.column_dimensions[L(i)].width = w
cell(u, "A1", "Юнит-экономика по товарам и каналам (месяц со средней сезонностью)", f_h1)
cell(u, "A2", "Логистика на 1 выкуп = (логистика до покупателя + обратная × доля невыкупа) / выкуп. Предельный ДРР — доля рекламы, при которой прибыль с единицы = 0. Серые столбцы справа — вспомогательные.", f_note)
header_row(u, 4, cols)
U0 = 5
r = U0
TOV = f"Товары!$B${G0}:$B${G1}"
def tov(col):  # подтянуть значение с листа «Товары» по SKU
    return f"INDEX(Товары!${col}${G0}:${col}${G1},MATCH($A{{r}},{TOV},0))"
def ch(col):   # подтянуть параметр канала с листа «Допущения»
    return f"INDEX(Допущения!${col}$6:${col}$8,MATCH($B{{r}},Допущения!$B$6:$B$8,0))"
share_col = {"Wildberries": "H", "Ozon": "I", "Сайт": "J"}
for sku, *_ in goods:
    for c in CH:
        cell(u, f"A{r}", sku); cell(u, f"B{r}", c)
        cell(u, f"C{r}", "=" + tov("C").format(r=r), f_link)
        cell(u, f"D{r}", "=" + tov("D").format(r=r), f_link)
        cell(u, f"T{r}", "=" + tov("E").format(r=r), f_link, RUB)                                     # базовая цена
        cell(u, f"U{r}", "=" + ch("C").format(r=r), f_link, PCT)                                     # комиссия %
        cell(u, f"V{r}", "=" + ch("D").format(r=r), f_link, RUB)                                     # логистика
        cell(u, f"W{r}", "=" + ch("E").format(r=r), f_link, RUB)                                     # обратная
        cell(u, f"P{r}", f"=MAX(0.05,MIN(1,{ch('F').format(r=r)}+Допущения!$C$25))", f_base, PCT)     # выкуп со сценарием
        cell(u, f"E{r}", f"=T{r}*(1+Допущения!$C$23)", f_base, RUB)                                 # цена со сценарием
        cell(u, f"F{r}", "=" + tov("F").format(r=r), f_link, RUB)
        cell(u, f"G{r}", f"=E{r}*U{r}", f_base, RUB)
        cell(u, f"H{r}", f"=(V{r}+W{r}*(1-P{r}))/P{r}", f_base, RUB)
        cell(u, f"I{r}", "=" + ch("H").format(r=r), f_link, RUB)
        cell(u, f"J{r}", f"=E{r}*{ch('G').format(r=r)}*Допущения!$C$24", f_base, RUB)
        cell(u, f"K{r}", f"=E{r}*Допущения!$C$11", f_base, RUB)
        cell(u, f"L{r}", f"=E{r}-F{r}-G{r}-H{r}-I{r}-J{r}-K{r}", f_b, RUB)
        cell(u, f"M{r}", f"=IFERROR(L{r}/E{r},0)", f_base, PCT)
        cell(u, f"N{r}", f"=IFERROR((E{r}-F{r}-G{r}-H{r}-I{r}-K{r})/E{r},0)", f_base, PCT)
        cell(u, f"O{r}", f"={tov('G').format(r=r)}*{tov(share_col[c]).format(r=r)}*(1+Допущения!$C$26)", f_base, NUM)
        cell(u, f"Q{r}", f"=O{r}*P{r}", f_base, NUM)
        cell(u, f"R{r}", f"=Q{r}*E{r}", f_base, RUB)
        cell(u, f"S{r}", f"=Q{r}*L{r}", f_b, RUB)
        for col in "TUVW":
            u[f"{col}{r}"].fill = fill_tot
        r += 1
U1 = r - 1
cell(u, f"C{r}", "Итого в месяц", f_b, fill=fill_tot, border=b_top)
for col in "OQRS":
    cell(u, f"{col}{r}", f"=SUM({col}{U0}:{col}{U1})", f_b, NUM if col in "OQ" else RUB, fill=fill_tot, border=b_top)
cell(u, f"M{r}", f"=IFERROR(S{r}/R{r},0)", f_b, PCT, fill=fill_tot, border=b_top)
UT = r
red = PatternFill("solid", fgColor="FBE3E3")
u.conditional_formatting.add(f"A{U0}:S{U1}", __import__("openpyxl").formatting.rule.FormulaRule(formula=[f"$L{U0}<0"], fill=red))
u.conditional_formatting.add(f"M{U0}:M{U1}", ColorScaleRule(start_type="num", start_value=-0.1, start_color="F8696B", mid_type="num", mid_value=0.1, mid_color="FFEB84", end_type="num", end_value=0.3, end_color="63BE7B"))
u.freeze_panes = "C5"
u.auto_filter.ref = f"A4:W{U1}"

# =====================================================================
# 5. P&L
# =====================================================================
p = wb.create_sheet("P&L")
p.column_dimensions["A"].width = 3; p.column_dimensions["B"].width = 34
for i in range(12):
    p.column_dimensions[L(3 + i)].width = 12
p.column_dimensions["O"].width = 14; p.column_dimensions["P"].width = 10
cell(p, "B2", "Отчёт о прибылях и убытках на 2027 год, ₽", f_h1)
cell(p, "B3", '="Сценарий: "&Допущения!C22', f_b)
header_row(p, 5, ["Статья"] + MONTHS + ["Итого за год", "% выручки"], start=2)
rng = lambda col: f"'Юнит-экономика'!${col}${U0}:${col}${U1}"
Q = rng("Q")
base_sums = {  # месячная база при средней сезонности
    "rev": f"SUMPRODUCT({Q},{rng('E')})", "cogs": f"SUMPRODUCT({Q},{rng('F')})", "com": f"SUMPRODUCT({Q},{rng('G')})",
    "log": f"SUMPRODUCT({Q},{rng('H')})+SUMPRODUCT({Q},{rng('I')})", "ads": f"SUMPRODUCT({Q},{rng('J')})", "tax": f"SUMPRODUCT({Q},{rng('K')})",
}
cell(p, "B6", "Индекс сезонности", f_note)
cell(p, "B7", "Фактор роста", f_note)
rows = [("Выручка", "rev", 8), ("Себестоимость", "cogs", 9), ("Комиссия площадок", "com", 10), ("Логистика и хранение", "log", 11),
        ("Реклама", "ads", 12), ("Налог", "tax", 13)]
for i in range(12):
    c = L(3 + i)
    cell(p, f"{c}6", f"=Допущения!{L(3 + i)}31", f_link, IDX)
    cell(p, f"{c}7", f"=(1+Допущения!$C$13)^{i}", f_base, IDX)
    for label, key, rr in rows:
        cell(p, f"{c}{rr}", f"={base_sums[key]}*{c}$6*{c}$7", f_base, RUB)
    cell(p, f"{c}14", f"={c}8-SUM({c}9:{c}13)", f_b, RUB, border=b_top)
    cell(p, f"{c}15", f"=Допущения!$C$12", f_link, RUB)
    cell(p, f"{c}16", f"={c}14-{c}15", f_b, RUB, border=b_top)
    cell(p, f"{c}17", f"=SUM($C$16:{c}16)", f_base, RUB)
    cell(p, f"{c}18", f"=IFERROR({c}14/{c}8,0)", f_base, PCT)
    cell(p, f"{c}19", f"=IFERROR({c}12/{c}8,0)", f_base, PCT)
for label, key, rr in rows:
    cell(p, f"B{rr}", label, f_b if key == "rev" else f_base)
cell(p, "B14", "Маржинальная прибыль", f_b, border=b_top)
cell(p, "B15", "Постоянные расходы", f_base)
cell(p, "B16", "Операционная прибыль", f_b, border=b_top)
cell(p, "B17", "Операционная прибыль нарастающим итогом", f_base)
cell(p, "B18", "Маржинальность, %", f_base)
cell(p, "B19", "ДРР, %", f_base)
for rr in range(8, 18):
    cell(p, f"O{rr}", f"=SUM(C{rr}:N{rr})", f_b, RUB, fill=fill_tot)
    if rr != 17:
        cell(p, f"P{rr}", f"=IFERROR(O{rr}/$O$8,0)", f_base, PCT)
p["O17"].value = "=N17"
cell(p, "O18", "=IFERROR(O14/O8,0)", f_b, PCT, fill=fill_tot)
cell(p, "O19", "=IFERROR(O12/O8,0)", f_b, PCT, fill=fill_tot)
for rr in (8, 14, 16):
    for i in range(14):
        p[f"{L(2 + i)}{rr}"].fill = fill_sub if rr != 16 else PatternFill("solid", fgColor="DDEBDD")
p.conditional_formatting.add("C16:N17", CellIsRule(operator="lessThan", formula=["0"], font=Font(name=F, color="C00000", bold=True)))
cell(p, "B20", "Месяцев с операционным убытком", f_base); cell(p, "C20", '=COUNTIF(C16:N16,"<0")', f_b, NUM)
cell(p, "B21", "Годовой коэффициент (сумма сезонности × рост)", f_note); cell(p, "C21", "=SUMPRODUCT(C6:N6,C7:N7)", f_note, IDX)

# итоги по каналам и категориям за год
cell(p, "B24", "Итоги за год по каналам", f_b)
header_row(p, 25, ["Канал", "Выручка, ₽", "Маржин. прибыль, ₽", "Маржин-сть", "Доля выручки"], start=2)
for k, c in enumerate(CH, start=26):
    cell(p, f"B{k}", c, f_b)
    cell(p, f"C{k}", f"=SUMIFS({rng('R')},{rng('B')},B{k})*$C$21", f_base, RUB)
    cell(p, f"D{k}", f"=SUMIFS({rng('S')},{rng('B')},B{k})*$C$21", f_base, RUB)
    cell(p, f"E{k}", f"=IFERROR(D{k}/C{k},0)", f_base, PCT)
    cell(p, f"F{k}", f"=IFERROR(C{k}/SUM($C$26:$C$28),0)", f_base, PCT)
cell(p, "B29", "Проверка: сумма = P&L", f_note); cell(p, "C29", '=IF(ABS(SUM(C26:C28)-O8)<1,"сходится","ОШИБКА")', f_note)

cell(p, "B31", "Итоги за год по категориям", f_b)
header_row(p, 32, ["Категория", "Выручка, ₽", "Маржин. прибыль, ₽", "Маржин-сть", "Доля выручки"], start=2)
cats = ["Текстиль", "Посуда", "Хранение", "Декор", "Ванная"]
for k, c in enumerate(cats, start=33):
    cell(p, f"B{k}", c, f_b)
    cell(p, f"C{k}", f"=SUMIFS({rng('R')},{rng('D')},B{k})*$C$21", f_base, RUB)
    cell(p, f"D{k}", f"=SUMIFS({rng('S')},{rng('D')},B{k})*$C$21", f_base, RUB)
    cell(p, f"E{k}", f"=IFERROR(D{k}/C{k},0)", f_base, PCT)
    cell(p, f"F{k}", f"=IFERROR(C{k}/SUM($C$33:$C$37),0)", f_base, PCT)
p.conditional_formatting.add("E26:E28", ColorScaleRule(start_type="min", start_color="F8CBAD", end_type="max", end_color="C6EFCE"))
p.conditional_formatting.add("E33:E37", ColorScaleRule(start_type="min", start_color="F8CBAD", end_type="max", end_color="C6EFCE"))
p.freeze_panes = "C6"

ch1 = BarChart(); ch1.type = "col"; ch1.title = "Операционная прибыль по месяцам, ₽"; ch1.style = 10
ch1.add_data(Reference(p, min_col=2, max_col=14, min_row=16), titles_from_data=True, from_rows=True)
ch1.set_categories(Reference(p, min_col=3, max_col=14, min_row=5))
ch1.y_axis.numFmt = '#,##0'; ch1.height = 7.5; ch1.width = 18; ch1.legend = None
ch1.y_axis.majorGridlines = None
p.add_chart(ch1, "H24")
ch2 = LineChart(); ch2.title = "Выручка и маржинальная прибыль, ₽"; ch2.style = 12
ch2.add_data(Reference(p, min_col=2, max_col=14, min_row=8), titles_from_data=True, from_rows=True)
ch2.add_data(Reference(p, min_col=2, max_col=14, min_row=14), titles_from_data=True, from_rows=True)
ch2.set_categories(Reference(p, min_col=3, max_col=14, min_row=5))
ch2.y_axis.numFmt = '#,##0'; ch2.height = 7.5; ch2.width = 18
p.add_chart(ch2, "H40")

# =====================================================================
# 6. Чувствительность
# =====================================================================
s = wb.create_sheet("Чувствительность")
s.column_dimensions["A"].width = 3; s.column_dimensions["B"].width = 22
for i in range(9):
    s.column_dimensions[L(3 + i)].width = 13
cell(s, "B2", "Годовая операционная прибыль: цена × доля рекламы, ₽", f_h1)
cell(s, "B3", "Строки — изменение цены от базовой, столбцы — ДРР во всех каналах. Объёмы продаж — из выбранного сценария, спрос от цены не зависит.", f_note)
cell(s, "B5", "Цена \\ ДРР", f_h2, fill=fill_head, al=Alignment(horizontal="center"))
drrs = [0.04, 0.06, 0.08, 0.10, 0.12, 0.14, 0.16, 0.18, 0.20]
dps = [-0.10, -0.05, 0.0, 0.05, 0.10, 0.15]
for j, d in enumerate(drrs):
    cell(s, f"{L(3 + j)}5", d, Font(name=F, size=10, bold=True, color="0000FF"), PCT, fill=fill_sub, al=Alignment(horizontal="center"))
for i, dp_ in enumerate(dps):
    rr = 6 + i
    cell(s, f"B{rr}", dp_, Font(name=F, size=10, bold=True, color="0000FF"), '+0%;-0%;0%', fill=fill_sub, al=Alignment(horizontal="center"))
    for j in range(len(drrs)):
        c = L(3 + j)
        price = f"{rng('T')}*(1+$B{rr})"
        f = (f"=SUMPRODUCT({Q},{price}*(1-{rng('U')}-Допущения!$C$11-{c}$5)-{rng('F')}-{rng('H')}-{rng('I')})"
             f"*'P&L'!$C$21-Допущения!$C$12*12")
        cell(s, f"{c}{rr}", f, f_base, RUB)
s.conditional_formatting.add("C6:K11", ColorScaleRule(start_type="min", start_color="F8696B", mid_type="num", mid_value=0, mid_color="FFFFFF", end_type="max", end_color="63BE7B"))
cell(s, "B13", "Как читать", f_b)
cell(s, "B14", "Зелёное — прибыль, красное — убыток. Сравните шаг по строке (+5% цены) и по столбцу (−2 п.п. рекламы): видно, какой рычаг сильнее.", f_base)
cell(s, "B16", "Прибыль +5% цены при текущем ДРР 12%", f_base); cell(s, "F16", "=G9-G8", f_b, RUB)
cell(s, "B17", "Прибыль при снижении ДРР с 12% до 8%", f_base); cell(s, "F17", "=E8-G8", f_b, RUB)

for wsx in wb.worksheets:
    wsx.sheet_view.showGridLines = wsx.title in ("Юнит-экономика",)
wb.save(OUT)
print("saved", OUT)
